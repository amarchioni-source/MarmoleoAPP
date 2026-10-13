"""Compara una planilla generada con una planilla de referencia armada a mano.

Uso:
    python verificar.py listing.xlsx teins.xlsx serie_inicial referencia.xlsx
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook

from marmoleo_core import generar


def datos(wb):
    out = {}
    for ws in wb.worksheets:
        filas, r = [], 12
        while isinstance(ws.cell(r, 2).value, (int, float)):
            filas.append((ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value))
            r += 1
        out[ws.title] = dict(serie=ws["N2"].value, tropa=ws["N6"].value,
                             fecha=ws["M7"].value.date(), filas=filas, fila_total=r,
                             formula=ws.cell(r, 5).value)
    return out


def main(listing, teins, serie, ref):
    with open(listing, "rb") as a, open(teins, "rb") as b:
        res = generar(a, b, serie)
    gen = datos(load_workbook(io.BytesIO(res.contenido)))
    exp = datos(load_workbook(ref))
    ok = True
    if list(gen) != list(exp):
        ok = False
        print("Hojas distintas:", list(gen), "vs", list(exp))
    for k in exp:
        for campo in ("serie", "tropa", "fecha", "filas", "fila_total"):
            if gen.get(k, {}).get(campo) != exp[k][campo]:
                ok = False
                print(f"[{k}] difiere {campo}: generado={gen.get(k, {}).get(campo)!r:.80} esperado={exp[k][campo]!r:.80}")
    print("Archivo:", res.nombre_archivo, "| hojas:", len(res.hojas))
    print("Advertencias:", res.advertencias or "ninguna")
    print("RESULTADO:", "IGUAL a la referencia" if ok else "HAY DIFERENCIAS")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    if len(sys.argv) != 5:
        sys.exit(__doc__)
    main(*sys.argv[1:])
