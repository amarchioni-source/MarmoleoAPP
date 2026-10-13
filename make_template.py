"""Genera plantilla_base.xlsx a partir de una planilla de marmoleo ya armada.

Uso:
    python make_template.py planilla_marmoleo_0910.xlsx

La plantilla resultante tiene dos hojas:
  - BASE:    encabezado de la planilla (filas 1 a 11), sin datos variables, con el logo.
  - ESTILOS: una fila modelo de datos (fila 1) y una fila modelo de total (fila 2),
             de las que la app copia el formato.
Solo hace falta correrlo si cambia el diseño de la planilla.
"""
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

HEADER_ROWS = 11
DATA_PROTO_ROW = 20     # fila de datos "estándar" en la planilla de origen
TOTAL_ROW_GUESS = None  # se detecta buscando "TOTAL AUTORIZADAS"
LAST_COL = 15           # A..O


def main(src: str) -> None:
    wb = load_workbook(src)
    ws = wb[wb.sheetnames[0]]

    # Fila de total: la que contiene "TOTAL AUTORIZADAS"
    total_row = next(
        r for r in range(HEADER_ROWS + 1, ws.max_row + 1)
        if any(isinstance(c.value, str) and "TOTAL AUTORIZADAS" in c.value for c in ws[r])
    )

    # Hoja de estilos con las filas modelo
    est = wb.create_sheet("ESTILOS")
    for c in range(1, LAST_COL + 1):
        for dst_row, src_row in ((1, DATA_PROTO_ROW), (2, total_row)):
            s = ws.cell(src_row, c)
            d = est.cell(dst_row, c)
            d.font, d.border, d.fill = copy(s.font), copy(s.border), copy(s.fill)
            d.alignment, d.number_format = copy(s.alignment), s.number_format

    # Limpiar la hoja base: quitar todo lo que está debajo del encabezado
    for rng in [str(m) for m in ws.merged_cells.ranges if m.min_row > HEADER_ROWS]:
        ws.unmerge_cells(rng)
    ws.delete_rows(HEADER_ROWS + 1, ws.max_row)

    # Vaciar los datos variables del encabezado (se conserva el formato)
    ws["N2"].value = None   # N° de serie
    ws["N6"].value = None   # N° de tropa
    ws["M7"].value = None   # fecha

    ws.title = "BASE"
    for name in wb.sheetnames:
        if name not in ("BASE", "ESTILOS"):
            del wb[name]

    out = Path(__file__).resolve().parent / "plantilla_base.xlsx"
    wb.save(out)
    print("Plantilla guardada en", out)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
