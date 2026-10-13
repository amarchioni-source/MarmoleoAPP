"""Lógica de armado de la planilla de marmoleo.

Entradas
  - Listing Faena (GD_informe_ListingFaena_*.xlsx): define qué tropas son W y en qué orden se faenaron.
  - Teins (DESCARGA_TEINS_*.xlsx): detalle animal por animal (garrón, clasificación, tipificación, amparo).

Reglas
  - Una hoja por tropa W (según el amparo del Listing), en el orden del Listing.
  - Dentro de cada hoja, solo los animales con amparo W en el Teins, ordenados por número de garrón.
  - Fecha de la planilla: día siguiente a la faena (o la que se indique).
  - N° de serie: correlativo por hoja, a partir del número inicial indicado.
"""
from __future__ import annotations

import io
import re
from collections import defaultdict
from copy import copy, deepcopy
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = BASE_DIR / "plantilla_base.xlsx"

FIRST_DATA_ROW = 12
LAST_COL = 15          # columna O
SERIE_MIN_WIDTH = 8    # '00018335'


class DatosInvalidos(Exception):
    """Error de contenido de los archivos, con mensaje apto para mostrar al usuario."""


# --------------------------------------------------------------------------- lectura

@dataclass
class TropaListing:
    tropa: int
    desde: int
    hasta: int
    cantidad: int
    amparo: str


@dataclass
class Listing:
    fecha: date | None
    tramos: list[TropaListing]


@dataclass
class Animal:
    tropa: int
    tip: str
    garron: int
    amparo: str
    clas: str


@dataclass
class Teins:
    fecha: date | None
    animales: list[Animal]


def _norm(v) -> str:
    return str(v).strip().lower() if v is not None else ""


def _as_int(v) -> int | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return None


def _as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", v)
        if m:
            d, mo, y = map(int, m.groups())
            try:
                return date(y, mo, d)
            except ValueError:
                return None
    return None


def _abrir(archivo, nombre: str):
    try:
        return load_workbook(archivo, data_only=True)
    except Exception as exc:  # archivo corrupto, no es xlsx, etc.
        raise DatosInvalidos(
            f"No pude abrir «{nombre}». Verifique que sea un archivo Excel (.xlsx) descargado del sistema."
        ) from exc


def leer_listing(archivo, nombre: str = "Listing Faena") -> Listing:
    ws = _abrir(archivo, nombre).worksheets[0]

    # Encabezado: la fila con "Tropa", "Garron Desde" y "Amparo"
    cols: dict[str, int] = {}
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=40):
        names = {_norm(c.value): c.column for c in row if c.value is not None}
        if {"tropa", "garron desde", "amparo"} <= names.keys():
            cols, header_row = names, row[0].row
            break
    if header_row is None:
        raise DatosInvalidos(
            f"«{nombre}» no tiene la tabla del Listing Faena (columnas Tropa, Garron Desde, Amparo). "
            "¿Se cargó el archivo correcto en este casillero?"
        )

    tramos: list[TropaListing] = []
    for r in range(header_row + 1, ws.max_row + 1):
        tropa = _as_int(ws.cell(r, cols["tropa"]).value)
        if tropa is None:
            continue
        desde = _as_int(ws.cell(r, cols["garron desde"]).value)
        hasta = _as_int(ws.cell(r, cols.get("garron hasta", cols["garron desde"])).value)
        if desde is None or hasta is None:
            continue
        cant = _as_int(ws.cell(r, cols["cant animales"]).value) if "cant animales" in cols else None
        amparo = str(ws.cell(r, cols["amparo"]).value or "").strip().upper()
        tramos.append(TropaListing(tropa, desde, hasta, cant if cant is not None else hasta - desde + 1, amparo))

    if not tramos:
        raise DatosInvalidos(f"«{nombre}» no tiene tropas cargadas.")

    # Fecha: la celda a la derecha de "Fecha:"
    fecha = None
    for row in ws.iter_rows(min_row=1, max_row=header_row):
        for i, c in enumerate(row):
            if isinstance(c.value, str) and c.value.strip().lower().startswith("fecha"):
                for d in row[i + 1:]:
                    fecha = _as_date(d.value)
                    if fecha:
                        break
            if fecha:
                break
        if fecha:
            break
    return Listing(fecha, tramos)


def leer_teins(archivo, nombre: str = "Teins") -> Teins:
    ws = _abrir(archivo, nombre).worksheets[0]

    cols: dict[str, int] = {}
    fecha = None
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        first = row[0].value
        if fecha is None and isinstance(first, str) and "FECHAFAENAD" in first.upper():
            fecha = _as_date(first)
        if not cols:
            names = {_norm(c.value): c.column for c in row if c.value is not None}
            if {"tropa", "gron", "amparo", "clas", "tip"} <= names.keys():
                cols = names
    if not cols:
        raise DatosInvalidos(
            f"«{nombre}» no tiene el detalle de faena (columnas TROPA, TIP, GRON, AMPARO, CLAS). "
            "¿Se cargó el archivo correcto en este casillero?"
        )

    animales: list[Animal] = []
    for r in range(1, ws.max_row + 1):
        tropa = _as_int(ws.cell(r, cols["tropa"]).value)
        garron = _as_int(ws.cell(r, cols["gron"]).value)
        if not tropa or garron is None:   # encabezados, títulos y filas sin tropa
            continue
        animales.append(Animal(
            tropa=tropa,
            tip=str(ws.cell(r, cols["tip"]).value or "").strip(),
            garron=garron,
            amparo=str(ws.cell(r, cols["amparo"]).value or "").strip().upper(),
            clas=str(ws.cell(r, cols["clas"]).value or "").strip(),
        ))
    if not animales:
        raise DatosInvalidos(f"«{nombre}» no tiene animales cargados.")
    return Teins(fecha, animales)


# --------------------------------------------------------------------------- armado

@dataclass
class HojaResumen:
    tropa: int
    animales: int
    serie: str
    garron_desde: int
    garron_hasta: int
    excluidos: dict[str, int] = field(default_factory=dict)


@dataclass
class Resultado:
    contenido: bytes
    nombre_archivo: str
    fecha_faena: date | None
    fecha_planilla: date
    hojas: list[HojaResumen]
    advertencias: list[str]
    faltantes: list[dict] = field(default_factory=list)


def formatear_serie(inicial: str, offset: int) -> str:
    texto = (inicial or "").strip()
    if not texto.isdigit():
        raise DatosInvalidos("El número de serie inicial debe tener solo dígitos (por ejemplo 00018335).")
    ancho = max(SERIE_MIN_WIDTH, len(texto))
    return str(int(texto) + offset).zfill(ancho)


def _tropas_w(listing: Listing) -> list[tuple[int, list[TropaListing]]]:
    """Tropas cuyo amparo en el Listing es W, en orden de faena (garrón inicial)."""
    por_tropa: dict[int, list[TropaListing]] = defaultdict(list)
    for t in listing.tramos:
        if t.amparo == "W":
            por_tropa[t.tropa].append(t)
    return sorted(por_tropa.items(), key=lambda kv: min(t.desde for t in kv[1]))


LOGO_PATH = BASE_DIR / "logo.jpeg"


def _logo(anchor) -> XLImage:
    """Logo de la planilla, con la misma posición y tamaño que en el modelo."""
    img = XLImage(io.BytesIO(LOGO_PATH.read_bytes()))
    img.anchor = deepcopy(anchor)
    return img


def armar_planilla(listing: Listing, teins: Teins, serie_inicial: str,
                   fecha_planilla: date | None = None) -> Resultado:
    advertencias: list[str] = []
    faltantes: list[dict] = []   # tropas W del Listing que no están (o no están completas) en el Teins

    # Fecha
    fecha_faena = listing.fecha or teins.fecha
    if listing.fecha and teins.fecha and listing.fecha != teins.fecha:
        advertencias.append(
            f"La fecha de faena no coincide: el Listing indica {listing.fecha:%d/%m/%Y} "
            f"y el Teins {teins.fecha:%d/%m/%Y}. Se usó la del Listing."
        )
    if fecha_planilla is None:
        if fecha_faena is None:
            raise DatosInvalidos("No encontré la fecha de faena en los archivos. Indique la fecha de la planilla.")
        fecha_planilla = fecha_faena + timedelta(days=1)

    serie_inicial = (serie_inicial or "").strip()
    formatear_serie(serie_inicial, 0)  # valida

    tropas_w = _tropas_w(listing)
    if not tropas_w:
        raise DatosInvalidos("El Listing Faena no tiene tropas con amparo W.")

    por_tropa: dict[int, list[Animal]] = defaultdict(list)
    for a in teins.animales:
        por_tropa[a.tropa].append(a)

    # Animales W en tropas que el Listing no marca como W
    w_set = {t for t, _ in tropas_w}
    for t, lista in sorted(por_tropa.items()):
        n = sum(1 for a in lista if a.amparo == "W")
        if n and t not in w_set:
            advertencias.append(
                f"Tropa {t}: el Teins tiene {n} animales con amparo W, pero el Listing no la marca como W. No se incluyó."
            )

    # Plantilla
    base_wb = load_workbook(TEMPLATE_PATH)
    base, estilos = base_wb["BASE"], base_wb["ESTILOS"]
    logo_anchor = base._images[0].anchor
    wb = base_wb

    hojas: list[HojaResumen] = []
    nuevas = []
    for idx, (tropa, tramos) in enumerate(tropas_w):
        todos = por_tropa.get(tropa, [])
        w = sorted((a for a in todos if a.amparo == "W"), key=lambda a: a.garron)
        excl: dict[str, int] = defaultdict(int)
        for a in todos:
            if a.amparo != "W":
                excl[a.amparo or "(sin amparo)"] += 1

        # Controles: el Listing contra el Teins
        esperados = sum(t.cantidad for t in tramos)
        desde, hasta = min(t.desde for t in tramos), max(t.hasta for t in tramos)
        if not todos:
            faltantes.append(dict(tropa=tropa, estado="ausente", esperados=esperados,
                                  desde=desde, hasta=hasta, garrones_faltantes=[]))
            continue
        if len(todos) < esperados:
            en_teins = {a.garron for a in todos}
            esperados_g = {g for t in tramos for g in range(t.desde, t.hasta + 1)}
            faltantes.append(dict(tropa=tropa, estado="incompleta", esperados=esperados,
                                  presentes=len(todos), desde=desde, hasta=hasta,
                                  garrones_faltantes=sorted(esperados_g - en_teins)))
        elif len(todos) > esperados:
            advertencias.append(
                f"Tropa {tropa}: el Listing indica {esperados} animales y el Teins tiene {len(todos)}."
            )
        if not w:
            advertencias.append(f"Tropa {tropa}: no hay animales con amparo W en el Teins. Se omitió la hoja.")
            continue
        fuera = [a.garron for a in w if not desde <= a.garron <= hasta]
        if fuera:
            advertencias.append(
                f"Tropa {tropa}: garrones fuera del rango del Listing ({desde}-{hasta}): {', '.join(map(str, fuera))}."
            )
        vistos, repetidos = set(), set()
        for a in w:
            (repetidos if a.garron in vistos else vistos).add(a.garron)
        if repetidos:
            advertencias.append(f"Tropa {tropa}: garrones repetidos: {', '.join(map(str, sorted(repetidos)))}.")
        sin_dato = [a.garron for a in w if not a.clas or not a.tip]
        if sin_dato:
            advertencias.append(
                f"Tropa {tropa}: garrones sin clasificación o tipificación: {', '.join(map(str, sin_dato))}."
            )

        serie = formatear_serie(serie_inicial, len(nuevas))
        ws = wb.copy_worksheet(base)
        ws.title = str(tropa)
        ws.add_image(_logo(logo_anchor))

        ws["N2"] = serie
        ws["N6"] = tropa
        ws["M7"] = datetime.combine(fecha_planilla, datetime.min.time())
        ws["M7"].number_format = "dd/mm/yyyy"  # fijo, para que no dependa de la configuración regional

        # Filas de datos
        r = FIRST_DATA_ROW
        for a in w:
            for c in range(1, LAST_COL + 1):
                ws.cell(r, c)._style = copy(estilos.cell(1, c)._style)
            ws.cell(r, 2, a.garron)
            ws.cell(r, 3, a.clas)
            ws.cell(r, 4, a.tip)
            r += 1

        # Fila de total
        total = r
        for c in range(1, LAST_COL + 1):
            ws.cell(total, c)._style = copy(estilos.cell(2, c)._style)
        for rng in (f"A{total}:D{total}", f"E{total}:F{total}", f"G{total}:J{total}", f"K{total}:O{total}"):
            ws.merge_cells(rng)
        ws.cell(total, 5, f"=SUM(E{FIRST_DATA_ROW}:E{total - 1})")
        ws.cell(total, 7, "TOTAL AUTORIZADAS")

        # Impresión: una hoja de ancho
        ws.print_area = f"A1:O{total}"
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        nuevas.append(ws)
        hojas.append(HojaResumen(tropa, len(w), serie, min(a.garron for a in w), max(a.garron for a in w),
                                 dict(excl)))

    if not nuevas:
        if faltantes:
            lista = ", ".join(str(f["tropa"]) for f in faltantes)
            raise DatosInvalidos(
                f"Ninguna tropa W del Listing figura completa en el Teins (tropas: {lista}). "
                "Posiblemente aún no estén valorizadas: corresponde reclamar a Compra de Hacienda."
            )
        raise DatosInvalidos("No se generó ninguna hoja: no hay animales W en las tropas W del Listing.")

    del wb["BASE"], wb["ESTILOS"]
    wb.active = 0
    for i, ws in enumerate(wb.worksheets):
        ws.sheet_view.tabSelected = i == 0
    wb.calculation.fullCalcOnLoad = True

    buf = io.BytesIO()
    wb.save(buf)
    return Resultado(
        contenido=buf.getvalue(),
        nombre_archivo=f"planilla_marmoleo_{fecha_planilla:%d%m}.xlsx",
        fecha_faena=fecha_faena,
        fecha_planilla=fecha_planilla,
        hojas=hojas,
        advertencias=advertencias,
        faltantes=faltantes,
    )


def generar(listing_file, teins_file, serie_inicial: str, fecha_planilla: date | None = None,
            nombre_listing: str = "Listing Faena", nombre_teins: str = "Teins") -> Resultado:
    return armar_planilla(leer_listing(listing_file, nombre_listing),
                          leer_teins(teins_file, nombre_teins),
                          serie_inicial, fecha_planilla)
